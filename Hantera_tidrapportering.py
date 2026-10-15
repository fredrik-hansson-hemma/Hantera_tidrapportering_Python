"""
Tidrapportering - Python-konvertering från SAS
Läser tidrapporter från Grindstone och skapar rapporter för Antura och flextid
"""

import sys
from datetime import datetime, timedelta

import openpyxl
import polars as pl
from dateutil.relativedelta import relativedelta
from openpyxl.styles import Font, PatternFill

from config import (
    ANTURA_EXCEL_FILE,
    CATEGORY_HTML_FILE,
    DATA_DIR,
    FLEX_EXCEL_FILE,
    GRINDSTONE_FILE,
)
from filhjalp import save_with_retry
from grindstone_reader import load_grindstone_data
from weekly_category_report import create_category_report

MANADSNAMN = [
    "", "Januari", "Februari", "Mars", "April", "Maj", "Juni",
    "Juli", "Augusti", "September", "Oktober", "November", "December",
]

# ============================================================================
# FUNKTIONER
# ============================================================================


def create_all_days(df: pl.DataFrame) -> pl.DataFrame:
    """
    Skapar upp alla dagar för att fylla ut dataset (för lediga dagar)
    Motsvarar data alla_dagar i SAS
    """
    min_datum = df.select(pl.col("startdatum").min()).item()
    max_datum = df.select(pl.col("startdatum").max()).item()

    # Första och sista dagen i månaden
    forsta = min_datum.replace(day=1)
    sista = (max_datum.replace(day=1) + relativedelta(months=1)) - timedelta(days=1)

    date_range = pl.date_range(forsta, sista, interval="1d", eager=True)

    return pl.DataFrame({"startdatum": date_range, "Duration": [0] * len(date_range)})


def create_main_query(df: pl.DataFrame, all_days: pl.DataFrame) -> pl.DataFrame:
    """
    Skapar huvuddataset med alla beräknade kolumner
    Motsvarar WORK.QUERY_FOR_THISMONTH_DERIVED_COLU i SAS
    """
    # Full outer join mellan tiddata och alla dagar
    merged = df.join(all_days, on="startdatum", how="full", suffix="_all")

    merged = merged.with_columns(
        [
            # Duration - använd från tiddata om den finns, annars från all_days
            pl.coalesce(["Duration", "Duration_all"]).alias("Duration"),
            # DurationTimmar - konvertera sekunder till timmar
            (pl.coalesce(["Duration", "Duration_all"]) / 3600.0).alias("DurationTimmar"),
            pl.coalesce(["startdatum", "startdatum_all"]).alias("startdatum"),
        ]
    )

    # Vecka är ISO-vecka och hör ihop med ISO-år (iso_ar). Kalenderåret (ar)
    # används tillsammans med manad. Vid årsskiftet skiljer de sig åt, t.ex.
    # 2025-12-29 är vecka 1 år 2026.
    merged = merged.with_columns(
        [
            pl.col("startdatum").dt.week().alias("Vecka"),
            pl.col("startdatum").dt.iso_year().alias("iso_ar"),
            pl.col("startdatum").dt.strftime("%a").alias("veckodag"),
            pl.col("startdatum").dt.month().alias("manad"),
            pl.col("startdatum").dt.year().alias("ar"),
        ]
    )

    cols = [
        "Work Item",
        "Start",
        "End",
        "Duration",
        "DurationTimmar",
        "startdatum",
        "Vecka",
        "iso_ar",
        "veckodag",
        "manad",
        "ar",
        "Notes",
    ]

    return merged.select([c for c in cols if c in merged.columns])


def add_time_type(df: pl.DataFrame) -> pl.DataFrame:
    """
    Lägger till kategori för typ av tid
    Motsvarar WORK.typ_av_tid i SAS
    """
    twelve_months_ago = datetime.now().date() - timedelta(days=365)

    return df.filter(pl.col("startdatum") > twelve_months_ago).with_columns(
        pl.when(~pl.col("Work Item").str.starts_with("Från"))
        .then(pl.lit("Arbetstid"))
        .otherwise(pl.col("Work Item"))
        .alias("Typ_av_tid")
    )


def create_antura_report(df: pl.DataFrame, output_file):
    """
    Skapar Antura-rapport i Excel
    Motsvarar PROC TABULATE för Antura i SAS
    """
    # Filtrera till förra månaden och framåt
    first_day_last_month = datetime.now().date().replace(day=1) - relativedelta(months=1)
    df_filtered = df.filter(pl.col("startdatum") >= first_day_last_month)

    wb = openpyxl.Workbook()
    wb.remove(wb.active)

    weeks = df_filtered.select(["iso_ar", "Vecka"]).unique().sort(["iso_ar", "Vecka"])

    for iso_ar, vecka in weeks.iter_rows():
        ws = wb.create_sheet(title=f"Vecka {vecka}")

        week_data = df_filtered.filter(
            (pl.col("iso_ar") == iso_ar) & (pl.col("Vecka") == vecka)
        )

        pivot_df = week_data.sql("""
            SELECT
                "Work Item",
                SUM(CASE WHEN veckodag = 'Mon' THEN DurationTimmar ELSE 0 END) as Mån,
                SUM(CASE WHEN veckodag = 'Tue' THEN DurationTimmar ELSE 0 END) as Tis,
                SUM(CASE WHEN veckodag = 'Wed' THEN DurationTimmar ELSE 0 END) as Ons,
                SUM(CASE WHEN veckodag = 'Thu' THEN DurationTimmar ELSE 0 END) as Tor,
                SUM(CASE WHEN veckodag = 'Fri' THEN DurationTimmar ELSE 0 END) as Fre,
                SUM(CASE WHEN veckodag = 'Sat' THEN DurationTimmar ELSE 0 END) as Lör,
                SUM(CASE WHEN veckodag = 'Sun' THEN DurationTimmar ELSE 0 END) as Sön,
                SUM(DurationTimmar) as Veckosumma
            FROM self
            WHERE "Work Item" IS NOT NULL
            GROUP BY "Work Item"
            ORDER BY "Work Item"
        """)
        pivot_dict = pivot_df.to_dicts()

        day_columns = ["Mån", "Tis", "Ons", "Tor", "Fre", "Lör", "Sön", "Veckosumma"]
        day_totals = ["Dagsumma"] + [
            sum(row.get(col, 0) or 0 for row in pivot_dict) for col in day_columns
        ]

        # Rubrikraden läses av antura_filler.py för att veckorna ska kunna
        # sorteras rätt även över årsskiftet
        first_day = week_data.select(pl.col("startdatum").min()).item()
        last_day = week_data.select(pl.col("startdatum").max()).item()
        ws.append([f"Vecka {vecka}: {first_day} - {last_day}"])
        ws.append([])

        ws.append(["Work Item"] + day_columns)
        for row_dict in pivot_dict:
            ws.append(
                [row_dict.get("Work Item", "")]
                + [row_dict.get(col, 0) for col in day_columns]
            )
        ws.append(day_totals)

        # Formatera
        for cell in ws[3]:
            cell.font = Font(bold=True)
            cell.fill = PatternFill(start_color="CCCCCC", fill_type="solid")

        for row in ws.iter_rows(min_row=4, min_col=2):
            for cell in row:
                if isinstance(cell.value, (int, float)):
                    cell.number_format = "#,##0.00"

        # Anteckningar – en rad per dag, work item och anteckningstext.
        # Läses av antura_filler.py som listar dem efter ifyllnad.
        notes_data = (
            week_data.filter(
                pl.col("Notes").is_not_null() & (pl.col("Notes").str.strip_chars() != "")
            )
            .group_by(["startdatum", "Work Item", "Notes"])
            .agg(pl.col("DurationTimmar").sum())
            .sort(["startdatum", "Work Item"])
        )

        if len(notes_data) > 0:
            ws.append([])
            ws.append(["Anteckningar"])
            ws.append(["Datum", "Work Item", "Anteckning", "Timmar"])

            for row in notes_data.select(
                ["startdatum", "Work Item", "Notes", "DurationTimmar"]
            ).iter_rows():
                ws.append(list(row))

    save_with_retry(lambda: wb.save(output_file), output_file)
    print(f"Antura-rapport sparad: {output_file}")


def seconds_to_hhmm(seconds) -> str:
    """Konverterar sekunder till H:MM, avrundat till närmaste minut."""
    if not seconds:
        return ""
    rounded_seconds = int(seconds) + 30
    hours = rounded_seconds // 3600
    minutes = (rounded_seconds % 3600) // 60
    return f"{hours}:{minutes:02d}"


def create_flex_report(df: pl.DataFrame, output_file):
    """
    Skapar flextidsrapport i Excel med tid i HH:MM format
    Motsvarar PROC TABULATE för flex i SAS
    """
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Flextid"

    months = df.select(["manad", "ar"]).unique().sort(["ar", "manad"])

    row_num = 1

    for month_row in months.iter_rows(named=True):
        ar = month_row["ar"]
        manad = month_row["manad"]

        month_data = df.filter((pl.col("ar") == ar) & (pl.col("manad") == manad))

        pivot_df = month_data.sql("""
            SELECT
                startdatum,
                veckodag,
                MAX(CASE WHEN Typ_av_tid = 'Arbetstid' THEN Duration ELSE NULL END) as Arbetstid,
                MAX(CASE WHEN Typ_av_tid LIKE 'Från%' THEN Duration ELSE NULL END) as Frånvaro,
                SUM(Duration) as Total
            FROM self
            GROUP BY startdatum, veckodag
            ORDER BY startdatum
        """)

        # Månadstitel
        ws.cell(row_num, 1, f"{MANADSNAMN[manad]} {ar}")
        ws.cell(row_num, 1).font = Font(bold=True, size=14)
        row_num += 2

        headers = ["Datum", "Veckodag", "Arbetstid", "Frånvaro", "Total"]
        for col_num, header in enumerate(headers, 1):
            cell = ws.cell(row_num, col_num, header)
            cell.font = Font(bold=True)
            cell.fill = PatternFill(start_color="CCCCCC", fill_type="solid")
        row_num += 1

        total_arbetstid = 0
        total_franvaro = 0
        total_total = 0

        for data_row in pivot_df.iter_rows(named=True):
            ws.cell(row_num, 1, data_row["startdatum"])
            ws.cell(row_num, 2, data_row["veckodag"])
            ws.cell(row_num, 3, seconds_to_hhmm(data_row["Arbetstid"]))
            ws.cell(row_num, 4, seconds_to_hhmm(data_row["Frånvaro"]))
            ws.cell(row_num, 5, seconds_to_hhmm(data_row["Total"]))

            total_arbetstid += data_row["Arbetstid"] or 0
            total_franvaro += data_row["Frånvaro"] or 0
            total_total += data_row["Total"] or 0

            row_num += 1

        # Total rad
        ws.cell(row_num, 1, "Total (ALL)")
        ws.cell(row_num, 1).font = Font(bold=True)
        ws.cell(row_num, 3, seconds_to_hhmm(total_arbetstid))
        ws.cell(row_num, 4, seconds_to_hhmm(total_franvaro))
        ws.cell(row_num, 5, seconds_to_hhmm(total_total))

        for col in range(3, 6):
            ws.cell(row_num, col).font = Font(bold=True)

        row_num += 3  # Mellanrum till nästa månad

    for col, width in zip("ABCDE", [12, 10, 12, 12, 12]):
        ws.column_dimensions[col].width = width

    save_with_retry(lambda: wb.save(output_file), output_file)
    print(f"Flextidsrapport sparad: {output_file}")


# ============================================================================
# HUVUDPROGRAM
# ============================================================================


def main():
    print("=== Tidrapportering - Python version ===\n")

    DATA_DIR.mkdir(exist_ok=True)

    print(f"Läser Grindstone-databas: {GRINDSTONE_FILE}")
    try:
        df = load_grindstone_data(GRINDSTONE_FILE)
    except FileNotFoundError:
        print(f"ERROR: Grindstone-filen saknas: {GRINDSTONE_FILE}")
        print("Kontrollera sökvägen i config.py")
        sys.exit(1)
    print(f"  Läste {len(df)} rader")

    df = df.with_columns(pl.col("Start").dt.date().alias("startdatum"))

    print("\nSkapar huvuddataset...")
    main_df = create_main_query(df, create_all_days(df))

    print("\nSkapar Antura-rapport...")
    create_antura_report(main_df, ANTURA_EXCEL_FILE)

    print("\nSkapar flextidsrapport...")
    create_flex_report(add_time_type(main_df), FLEX_EXCEL_FILE)

    print("\nSkapar kategoriserad veckorapport (20 veckor)...")
    create_category_report(main_df, CATEGORY_HTML_FILE, weeks_back=20)

    print("\n=== Klart! ===")


if __name__ == "__main__":
    # Säkerställ att åäö och symboler kan skrivas ut oavsett konsolens teckentabell
    sys.stdout.reconfigure(encoding="utf-8")
    main()
